import csv, io, re
from openpyxl import load_workbook
from . import config

HEADERS = {'phone','phone_number','number','mobile','tel','전화번호','핸드폰','휴대폰','연락처','번호'}
def normalize(raw):
    if raw is None: return None
    s = str(raw).strip().replace('\u00a0','')
    if re.fullmatch(r'\d+\.0',s): s=s[:-2]
    s = re.sub(r'[\s()\-.]', '', s)
    if s.startswith('+82'): s = '0'+s[3:]
    elif s.startswith('82') and len(s) >= 11: s = '0'+s[2:]
    if not re.fullmatch(r'010\d{8}',s): return None
    if s.startswith('010') and len(s) != 11:return None
    return s

def display(phone):
    if len(phone)==11:return f'{phone[:3]}-{phone[3:7]}-{phone[7:]}'
    return f'{phone[:3]}-{phone[3:6]}-{phone[6:]}'

def _csv_rows(content):
    for enc in ('utf-8-sig','cp949','euc-kr'):
        try:
            text=content.decode(enc)
            sample=text[:8192]
            try: dialect=csv.Sniffer().sniff(sample,delimiters=',;\t')
            except csv.Error: dialect=csv.excel
            return list(csv.reader(io.StringIO(text),dialect))
        except UnicodeDecodeError:continue
    raise ValueError('CSV 인코딩을 해석할 수 없습니다.')

def parse_file(name, content):
    if len(content)>config.MAX_BYTES:raise ValueError('파일 용량 제한 초과')
    ext=name.lower().rsplit('.',1)[-1]
    if ext=='csv':rows=_csv_rows(content)
    elif ext=='txt':
        decoded=None
        for enc in ('utf-8-sig','cp949','euc-kr'):
            try:decoded=content.decode(enc);break
            except UnicodeDecodeError:pass
        if decoded is None:raise ValueError('TXT 인코딩을 해석할 수 없습니다.')
        rows=((line.strip(),) for line in decoded.splitlines() if line.strip())
    elif ext=='xlsx':
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        try: rows=wb.active.iter_rows(values_only=True)
        except Exception:wb.close();raise
    elif ext=='xls':
        try:import xlrd
        except ImportError:raise ValueError('XLS 지원을 위해 xlrd 설치가 필요합니다.')
        wb=xlrd.open_workbook(file_contents=content)
        sheet=wb.sheet_by_index(0)
        rows=(sheet.row_values(i) for i in range(sheet.nrows))
    else:raise ValueError('CSV/XLSX/XLS/TXT 형식만 허용합니다.')
    seen=set();phones=[];total=invalid=duplicates=0
    it=iter(rows);first=next(it,None)
    if first is None:raise ValueError('빈 파일입니다.')
    headers=[str(x or '').strip().lower() for x in first]
    col=next((i for i,h in enumerate(headers) if h in HEADERS),0)
    header=any(h in HEADERS for h in headers)
    def consume(row):
        nonlocal total,invalid,duplicates
        if total>=1_100_000:raise ValueError('파일 행 수가 안전 제한을 초과합니다.')
        total+=1
        phone=normalize(row[col] if col<len(row) else None)
        if not phone:invalid+=1
        elif phone in seen:duplicates+=1
        else:seen.add(phone);phones.append(phone)
    if not header:consume(first)
    for row in it:consume(row)
    if ext=='xlsx':wb.close()
    return {'total':total,'invalid':invalid,'duplicates':duplicates,'phones':phones,'valid':len(phones)}
